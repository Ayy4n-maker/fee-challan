from io import BytesIO
from datetime import date, datetime
import calendar

from flask import render_template, request, redirect, url_for, session, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app import app, get_portfolio_docs


def _admin_ok():
    return session.get("admin_logged_in") is True


def _value(v):
    if v is None:
        return ""
    if hasattr(v, "isoformat"):
        try:
            return v.isoformat()
        except Exception:
            pass
    if isinstance(v, dict):
        return "; ".join(f"{k}: {_value(x)}" for k, x in v.items())
    if isinstance(v, (list, tuple)):
        return ", ".join(_value(x) for x in v)
    return v


def _rows(collection):
    result = []
    for doc in get_portfolio_docs(collection):
        row = doc.to_dict() or {}
        row["document_id"] = doc.id
        result.append(row)
    return result


def _money(row, *names):
    for name in names:
        try:
            val = row.get(name)
            if val is not None and val != "":
                return float(val)
        except (TypeError, ValueError):
            pass
    return 0.0


def _student_rows():
    """Export students with the same ID format used on the website."""
    rows = _rows("students")
    result = []
    for index, row in enumerate(rows, start=1):
        serial = row.get("serial_id")
        try:
            display_id = "%03d" % int(serial) if serial is not None else "%03d" % index
        except (ValueError, TypeError):
            display_id = str(row.get("display_id") or row.get("id") or index)

        student = {
            "display_id": display_id,
            "student_name": row.get("student_name", ""),
            "father_name": row.get("father_name", ""),
            "roll_no": row.get("roll_no", ""),
            "class_name": row.get("class_name") or row.get("class", ""),
            "monthly_fee": _money(row, "monthly_fee", "tuition_fee"),
            "father_phone": row.get("father_phone") or row.get("student_phone") or row.get("parent_phone", ""),
            "mother_phone": row.get("mother_phone", ""),
            "date_of_birth": row.get("date_of_birth", ""),
            "gender": row.get("gender", ""),
            "address": row.get("address", ""),
            "previous_school": row.get("previous_school", ""),
            "previous_coaching": row.get("previous_coaching", ""),
        }
        result.append(student)
    return result


def _teacher_rows():
    """Export teachers with the same ID used on the website."""
    rows = _rows("teachers")
    result = []
    for index, row in enumerate(rows, start=1):
        display_id = row.get("display_id")
        if display_id in (None, ""):
            display_id = row.get("id")
        if display_id in (None, ""):
            display_id = index
        teacher = {
            "display_id": display_id,
            "teacher_name": row.get("teacher_name", ""),
            "allotted_class": row.get("allotted_class", ""),
            "allotted_subjects": row.get("allotted_subjects", ""),
            "phone": row.get("phone", ""),
            "salary": _money(row, "salary"),
        }
        result.append(teacher)
    return result


def _fee_payment_rows(records):
    """Export fee payment records with clean columns and formatted IDs."""
    result = []
    for index, row in enumerate(records, start=1):
        student_id = row.get("student_display_id") or row.get("student_id") or ""
        payment = {
            "student_display_id": student_id,
            "student_name": row.get("student_name", ""),
            "father_name": row.get("father_name", ""),
            "class_name": row.get("class_name") or row.get("class", ""),
            "fee_month": row.get("fee_month", ""),
            "monthly_fee": _money(row, "monthly_fee", "tuition_fee", "fee_amount"),
            "admission_fee": _money(row, "admission_fee", "additional_fee"),
            "annual_fee": _money(row, "annual_fee", "annual_charges"),
            "stationery_fee": _money(row, "stationery_fee", "exam_fee", "stationery_exam_fee"),
            "late_fee": _money(row, "late_fee", "late_charges"),
            "other_fee": _money(row, "other_fee", "other_charges"),
            "discount": _money(row, "discount"),
            "total_payable": _money(row, "total_payable", "total_amount"),
            "amount_received": _money(row, "amount_received", "amount_paid"),
            "status": str(row.get("status", "UNPAID")).upper(),
            "due_date": row.get("due_date", ""),
            "payment_date": row.get("payment_date", ""),
        }
        result.append(payment)
    return result


def _challan_rows(records):
    """Export challan records with clean columns."""
    result = []
    for index, row in enumerate(records, start=1):
        challan = {
            "challan_no": row.get("challan_no") or row.get("id") or f"{index:04d}",
            "student_display_id": row.get("student_display_id") or row.get("student_id", ""),
            "student_name": row.get("student_name", ""),
            "father_name": row.get("father_name", ""),
            "class_name": row.get("class_name") or row.get("class", ""),
            "fee_month": row.get("fee_month", ""),
            "total_payable": _money(row, "total_payable", "total_amount"),
            "due_date": row.get("due_date", ""),
            "status": str(row.get("status", "UNPAID")).upper(),
            "issue_date": row.get("issue_date") or row.get("challan_date", ""),
        }
        result.append(challan)
    return result


def _expense_rows(records):
    """Export expense records with clean columns."""
    result = []
    for index, row in enumerate(records, start=1):
        display_id = row.get("display_id") or row.get("serial_id") or index
        expense = {
            "display_id": display_id,
            "title": row.get("title") or row.get("expense_name", ""),
            "category": row.get("category", "General"),
            "amount": _money(row, "amount", "expense_amount"),
            "expense_date": row.get("expense_date") or row.get("date", ""),
            "description": row.get("description", ""),
        }
        result.append(expense)
    return result


HEADER_LABELS = {
    "display_id": "ID",
    "serial_id": "Serial No.",
    "id": "ID",
    "student_id": "Student ID",
    "student_display_id": "Student ID",
    "student_name": "Student Name",
    "father_name": "Father Name",
    "roll_no": "Roll No.",
    "class_name": "Class",
    "class": "Class",
    "monthly_fee": "Monthly Fee (Rs.)",
    "tuition_fee": "Tuition Fee (Rs.)",
    "admission_fee": "Admission Fee (Rs.)",
    "annual_fee": "Annual Charges (Rs.)",
    "stationery_fee": "Stationery / Exam (Rs.)",
    "late_fee": "Late Fee (Rs.)",
    "other_fee": "Other Charges (Rs.)",
    "discount": "Discount (Rs.)",
    "total_payable": "Total Payable (Rs.)",
    "amount_received": "Amount Received (Rs.)",
    "amount_paid": "Amount Paid (Rs.)",
    "amount": "Amount (Rs.)",
    "status": "Status",
    "fee_month": "Fee Month",
    "due_date": "Due Date",
    "payment_date": "Payment Date",
    "paid_at": "Paid At",
    "challan_no": "Challan No.",
    "challan_date": "Challan Date",
    "issue_date": "Issue Date",
    "date": "Date",
    "date_of_birth": "Date of Birth",
    "gender": "Gender",
    "phone": "Phone",
    "father_phone": "Father Phone",
    "mother_phone": "Mother Phone",
    "student_phone": "Student Phone",
    "parent_phone": "Parent Phone",
    "address": "Address",
    "previous_school": "Previous School",
    "previous_coaching": "Previous Coaching",
    "teacher_name": "Teacher Name",
    "allotted_class": "Allotted Class",
    "allotted_subjects": "Allotted Subjects",
    "salary": "Salary (Rs.)",
    "expense_name": "Expense Title",
    "title": "Expense Title",
    "category": "Category",
    "expense_date": "Expense Date",
    "description": "Description",
    "notes": "Notes",
    "source": "Revenue Source",
    "period": "Period",
    "revenue": "Total Revenue (Rs.)",
    "paid_payments": "Paid Records Count",
    "tuition": "Tuition Fees (Rs.)",
    "admission": "Admission Fees (Rs.)",
    "annual": "Annual Fees (Rs.)",
    "stationery_exam": "Stationery / Exam Fees (Rs.)",
    "late": "Late Fees (Rs.)",
    "other": "Other Fees (Rs.)",
    "discounts": "Discounts Given (Rs.)",
    "item": "Item",
}

# Border styles
_thin_border_side = Side(style="thin", color="CBD5E1")
thin_border = Border(
    left=_thin_border_side,
    right=_thin_border_side,
    top=_thin_border_side,
    bottom=_thin_border_side
)

_header_border_side = Side(style="thin", color="3B82F6")
header_border = Border(
    left=_header_border_side,
    right=_header_border_side,
    top=Side(style="medium", color="1E3A8A"),
    bottom=Side(style="medium", color="1E3A8A")
)


def _format_sheet_for_print(ws, is_landscape=True):
    """
    Configure print settings:
    - Fit all columns on 1 page wide
    - Rows can extend to unlimited pages (fitToHeight = 0)
    - Only print the boxes worked on (strict print_area, no blank gridlines)
    - Repeat header row on every printed page
    - Narrow margins for maximum readable space
    """
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.autoPageBreaks = True

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE if is_landscape else ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4

    # Narrow margins ensure columns fit cleanly with clear, readable fonts
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.page_margins.header = 0.3
    ws.page_margins.footer = 0.3

    # Repeat header row across all pages when rows extend to multiple pages
    if ws.max_row > 1:
        ws.print_title_rows = "1:1"

    # Only print the boxes that have been worked on (exact data boundary)
    max_row = max(ws.max_row, 1)
    max_col = max(ws.max_column, 1)
    max_col_letter = get_column_letter(max_col)
    ws.print_area = f"A1:{max_col_letter}{max_row}"

    # Do not print empty boxes/gridlines for unworked cells on the sheet
    ws.print_options.gridLines = False


def _make_sheet(wb, title, rows, is_landscape=True):
    ws = wb.create_sheet(title[:31])

    header_fill = PatternFill("solid", fgColor="1E3A8A")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    even_row_fill = PatternFill("solid", fgColor="F8FAFC")
    odd_row_fill = PatternFill("solid", fgColor="FFFFFF")
    data_font = Font(name="Calibri", size=10, color="1E293B")

    if not rows:
        ws.append(["No records found"])
        cell = ws.cell(row=1, column=1)
        cell.font = Font(name="Calibri", size=11, italic=True, color="64748B")
        cell.border = thin_border
        _format_sheet_for_print(ws, is_landscape=is_landscape)
        ws.column_dimensions["A"].width = 24
        return ws

    keys = []
    for row in rows:
        for key in row:
            if key not in keys:
                keys.append(key)

    # 1. Header row
    header_labels = [HEADER_LABELS.get(k, k.replace("_", " ").title()) for k in keys]
    ws.append(header_labels)
    ws.row_dimensions[1].height = 28

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border

    # 2. Data rows
    for row_idx, row in enumerate(rows, start=2):
        row_values = [_value(row.get(key)) for key in keys]
        ws.append(row_values)
        ws.row_dimensions[row_idx].height = 20
        fill = even_row_fill if row_idx % 2 == 0 else odd_row_fill

        for col_idx, key in enumerate(keys, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = data_font
            cell.fill = fill
            cell.border = thin_border

            # Smart alignment based on column type
            val = cell.value
            key_lower = key.lower()
            if isinstance(val, (int, float)) or any(m in key_lower for m in ("fee", "salary", "amount", "revenue", "tuition", "discount", "total")):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif any(c in key_lower for c in ("id", "roll", "date", "status", "gender", "period", "count")):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # 3. Dynamic readable column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if "\n" in val_str:
                val_str = max(val_str.split("\n"), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 32)

    # 4. Print setup
    _format_sheet_for_print(ws, is_landscape=is_landscape)
    return ws


def _date_text(row):
    for key in ("payment_date", "paid_at", "date", "expense_date", "created_at", "challan_date", "issue_date"):
        if row.get(key):
            return str(_value(row.get(key)))
    return "Unknown"


def _date_match(row, month, year):
    for key in ("payment_date", "paid_at", "date", "expense_date", "created_at", "challan_date", "issue_date"):
        value = row.get(key)
        if not value:
            continue
        text = str(value)
        if f"{year}-{month:02d}" in text or f"{month:02d}-{year}" in text:
            return True
    fee_month = str(row.get("fee_month") or "").lower()
    return calendar.month_name[month].lower() in fee_month or calendar.month_abbr[month].lower() in fee_month


def _build(month=None, year=None):
    monthly = month is not None and year is not None
    students = _student_rows()
    teachers = _teacher_rows()
    payments = _rows("fee_payments")
    expenses = _rows("expenses")
    challans = _rows("challans")

    report_payments = payments
    report_expenses = expenses
    report_challans = challans
    if monthly:
        report_payments = [p for p in payments if _date_match(p, month, year)]
        report_expenses = [e for e in expenses if _date_match(e, month, year)]
        report_challans = [c for c in challans if _date_match(c, month, year)]

    paid_payments = [p for p in report_payments if str(p.get("status", "")).upper() == "PAID"]

    revenue = sum(_money(p, "amount_received", "amount_paid", "total_payable", "total_amount") for p in paid_payments)
    tuition = sum(_money(p, "monthly_fee", "tuition_fee", "fee_amount") for p in paid_payments)
    late = sum(_money(p, "late_fee", "late_charges") for p in paid_payments)
    annual = sum(_money(p, "annual_fee", "annual_charges") for p in paid_payments)
    stationery = sum(_money(p, "stationery_fee", "exam_fee", "stationery_exam_fee") for p in paid_payments)
    admission = sum(_money(p, "admission_fee", "additional_fee") for p in paid_payments)
    other = sum(_money(p, "other_fee", "other_charges") for p in paid_payments)
    discount = sum(_money(p, "discount") for p in paid_payments)
    expenses_total = sum(_money(e, "amount", "expense_amount") for e in report_expenses)
    salaries = sum(_money(t, "salary") for t in teachers)
    total_expenses = expenses_total + salaries

    wb = Workbook()
    wb.remove(wb.active)

    # ----------------------------------------------------
    # Summary Sheet
    # ----------------------------------------------------
    summary = wb.create_sheet("Summary")

    # Title Banner
    summary.merge_cells("A1:B1")
    title_cell = summary["A1"]
    title_cell.value = "IESC BACKUP & FINANCIAL REPORT"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="1E3A8A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    summary.row_dimensions[1].height = 32

    # Info rows
    school_name = session.get("school_name", "Irshad Educational Academy")
    period_str = f"{calendar.month_name[month]} {year}" if monthly else "Complete history to date"
    gen_date = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    info_rows = [
        ("School", school_name),
        ("Report Period", period_str),
        ("Generated Date", gen_date),
    ]

    for label, val in info_rows:
        summary.append([label, val])
        r = summary.max_row
        summary.row_dimensions[r].height = 20
        c1, c2 = summary.cell(r, 1), summary.cell(r, 2)
        c1.font = Font(name="Calibri", size=10, bold=True, color="475569")
        c2.font = Font(name="Calibri", size=10, color="1E293B")
        c1.fill = PatternFill("solid", fgColor="F1F5F9")
        c2.fill = PatternFill("solid", fgColor="F8FAFC")
        c1.border = thin_border
        c2.border = thin_border
        c1.alignment = Alignment(horizontal="left", vertical="center")
        c2.alignment = Alignment(horizontal="left", vertical="center")

    summary.append([])

    # Financial Items
    summary.append(["Financial Item", "Amount (Rs.)"])
    fin_header_row = summary.max_row
    summary.row_dimensions[fin_header_row].height = 24
    for c in (summary.cell(fin_header_row, 1), summary.cell(fin_header_row, 2)):
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A8A")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    fin_items = [
        ("Total Revenue Received", revenue),
        ("Tuition / Monthly Fees", tuition),
        ("Admission / Additional Fees", admission),
        ("Annual Charges", annual),
        ("Stationery / Exam Fees", stationery),
        ("Late Charges", late),
        ("Other Charges", other),
        ("Discounts Given", discount),
        ("Recorded School Expenses", expenses_total),
        ("Teacher Salaries", salaries),
        ("Total School Expenses", total_expenses),
        ("Net Profit / Loss", revenue - total_expenses),
    ]

    for label, amount in fin_items:
        summary.append([label, amount])
        r = summary.max_row
        summary.row_dimensions[r].height = 20
        c1, c2 = summary.cell(r, 1), summary.cell(r, 2)
        is_highlight = label in ("Total Revenue Received", "Total School Expenses", "Net Profit / Loss")
        c1.font = Font(name="Calibri", size=10, bold=is_highlight, color="0F172A")
        c2.font = Font(
            name="Calibri",
            size=10,
            bold=is_highlight,
            color="15803D" if label == "Total Revenue Received" or (label == "Net Profit / Loss" and amount >= 0) else ("DC2626" if label in ("Total School Expenses") or (label == "Net Profit / Loss" and amount < 0) else "0F172A")
        )
        bg = "E2E8F0" if is_highlight else ("F8FAFC" if r % 2 == 0 else "FFFFFF")
        c1.fill = PatternFill("solid", fgColor=bg)
        c2.fill = PatternFill("solid", fgColor=bg)
        c1.border = thin_border
        c2.border = thin_border
        c1.alignment = Alignment(horizontal="left", vertical="center")
        c2.alignment = Alignment(horizontal="right", vertical="center")

    summary.append([])

    # Record Counts
    summary.append(["Record Type", "Total Records"])
    rec_header_row = summary.max_row
    summary.row_dimensions[rec_header_row].height = 24
    for c in (summary.cell(rec_header_row, 1), summary.cell(rec_header_row, 2)):
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E3A8A")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    rec_counts = [
        ("Registered Students", len(students)),
        ("Registered Teachers", len(teachers)),
        ("Fee Payment Records", len(report_payments)),
        ("Paid Fee Payments", len(paid_payments)),
        ("School Expenses", len(report_expenses)),
        ("Issued Challans", len(report_challans)),
    ]

    for label, count in rec_counts:
        summary.append([label, count])
        r = summary.max_row
        summary.row_dimensions[r].height = 20
        c1, c2 = summary.cell(r, 1), summary.cell(r, 2)
        c1.font = Font(name="Calibri", size=10, color="0F172A")
        c2.font = Font(name="Calibri", size=10, bold=True, color="1E3A8A")
        bg = "F8FAFC" if r % 2 == 0 else "FFFFFF"
        c1.fill = PatternFill("solid", fgColor=bg)
        c2.fill = PatternFill("solid", fgColor=bg)
        c1.border = thin_border
        c2.border = thin_border
        c1.alignment = Alignment(horizontal="left", vertical="center")
        c2.alignment = Alignment(horizontal="right", vertical="center")

    summary.column_dimensions["A"].width = 36
    summary.column_dimensions["B"].width = 24

    # Print setup for Summary: Portrait, fitToWidth 1, fitToHeight 0, print area A1:B{max_row}
    _format_sheet_for_print(summary, is_landscape=False)

    # ----------------------------------------------------
    # Detailed Data Sheets (Landscape for comfortable column fitting)
    # ----------------------------------------------------
    _make_sheet(wb, "Students", students, is_landscape=True)
    _make_sheet(wb, "Teachers", teachers, is_landscape=True)
    _make_sheet(wb, "Fee Payments", _fee_payment_rows(report_payments), is_landscape=True)
    _make_sheet(wb, "Challans", _challan_rows(report_challans), is_landscape=True)
    _make_sheet(wb, "Expenses", _expense_rows(report_expenses), is_landscape=True)

    _make_sheet(wb, "Revenue Sources", [
        {"source": "Tuition / Monthly Fees", "amount": tuition},
        {"source": "Admission / Additional Fees", "amount": admission},
        {"source": "Annual Charges", "amount": annual},
        {"source": "Stationery / Exam Fees", "amount": stationery},
        {"source": "Late Charges", "amount": late},
        {"source": "Other Charges", "amount": other},
        {"source": "Discounts Given", "amount": discount},
        {"source": "Total Revenue Received", "amount": revenue},
    ], is_landscape=False)

    _make_sheet(wb, "Monthly Revenue", [{
        "period": f"{calendar.month_name[month]} {year}" if monthly else "Complete history",
        "revenue": revenue,
        "paid_payments": len(paid_payments),
        "tuition": tuition,
        "admission": admission,
        "annual": annual,
        "stationery_exam": stationery,
        "late": late,
        "other": other,
        "discounts": discount,
    }], is_landscape=True)

    _make_sheet(wb, "Profit & Loss", [
        {"item": "Total Revenue", "amount": revenue},
        {"item": "Recorded Expenses", "amount": expenses_total},
        {"item": "Teacher Salaries", "amount": salaries},
        {"item": "Total School Expenses", "amount": total_expenses},
        {"item": "Net Profit / Loss", "amount": revenue - total_expenses},
    ], is_landscape=False)

    return wb


@app.route("/admin/backup")
def backup_page():
    if not _admin_ok():
        return redirect(url_for("admin_login"))
    today = date.today()
    return render_template(
        "backup.html",
        current_month=today.month,
        current_year=today.year,
        years=list(range(today.year - 5, today.year + 2)),
        month_names=list(calendar.month_name)[1:]
    )


@app.route("/admin/backup/download")
def download_backup():
    if not _admin_ok():
        return redirect(url_for("admin_login"))
    kind = request.args.get("backup_type", "complete")
    month = int(request.args.get("month", date.today().month)) if kind == "monthly" else None
    year = int(request.args.get("year", date.today().year)) if kind == "monthly" else None
    if kind == "monthly" and (month < 1 or month > 12 or year < 2000 or year > 2100):
        return redirect(url_for("backup_page"))

    wb = _build(month, year)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = (
        f"IESC_Backup_{year}_{month:02d}.xlsx"
        if kind == "monthly"
        else f"IESC_Complete_Backup_{date.today().isoformat()}.xlsx"
    )
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@app.after_request
def _add_backup_button(response):
    if request.path == "/admin" and response.status_code == 200 and response.content_type.startswith("text/html"):
        try:
            html = response.get_data(as_text=True)
            marker = "<!-- System -->"
            if marker in html and "/admin/backup" not in html:
                block = (
                    '<!-- Backup & Reports -->\n'
                    '        <div class="btn-group">\n'
                    '            <div class="group-label">💾 Backup & Reports</div>\n'
                    '            <div class="btn-row">\n'
                    '                <a href="/admin/backup" class="btn btn-teal">📥 Excel Backup & Reports</a>\n'
                    '            </div>\n'
                    '        </div>\n\n\n        '
                )
                response.set_data(html.replace(marker, block + marker, 1))
        except Exception as exc:
            print(f"Backup dashboard injection error: {exc}")
    return response


if __name__ == "__main__":
    app.run(debug=True)
